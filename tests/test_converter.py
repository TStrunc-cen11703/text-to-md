import asyncio
from io import BytesIO
import unittest
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import UploadFile
from pypdf import PdfWriter

from text_to_md.app import convert as convert_upload
from text_to_md.converter import _extract_comments, convert_docx, convert_pdf, convert_txt


class ConvertEndpointTests(unittest.TestCase):
    def test_uses_fixed_markdown_filename(self) -> None:
        document = Document()
        document.add_paragraph("Dokument pro test názvu výstupu.")
        output = BytesIO()
        document.save(output)
        upload = UploadFile(
            filename="jiny-nazev-vstupu.docx",
            file=BytesIO(output.getvalue()),
        )

        response = asyncio.run(convert_upload(upload))

        self.assertEqual(response.filename, "pozadavek-na-sablonu.md")


class ConvertDocxTests(unittest.TestCase):
    def make_document(self) -> bytes:
        document = Document()
        document.add_heading("Upozornění na neprovedené platby", level=1)
        document.add_heading("E-mail", level=2)
        document.add_paragraph("Předmět: <licenceName> - Upozornění")
        document.add_paragraph("Dobrý den, <licenceName>.")
        split_token = document.add_paragraph()
        split_token.add_run("<licence")
        split_token.add_run("Name>").bold = True
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Kanál"
        table.cell(0, 1).text = "E-mail"
        table.cell(1, 0).text = "Jazyk"
        table.cell(1, 1).text = "Čeština"
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def test_preserves_headings_text_and_tables(self) -> None:
        markdown, warnings = convert_docx(self.make_document(), "zadani.docx")

        self.assertIn("# Upozornění na neprovedené platby", markdown)
        self.assertIn("## Email", markdown)
        self.assertIn("**Předmět:** <licenceName> - Upozornění", markdown)
        self.assertIn("<licenceName>", markdown)
        self.assertIn("- **Kanál:** E-mail", markdown)
        self.assertIn("- **Jazyk:** Čeština", markdown)
        self.assertNotIn("Preheader", markdown)
        self.assertTrue(any("neodhadují" in warning for warning in warnings))

    def test_adds_fallback_title_when_word_has_no_title(self) -> None:
        document = Document()
        document.add_paragraph("Obyčejný odstavec.")
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "bez-nadpisu.docx")

        self.assertTrue(markdown.startswith("# bez-nadpisu\n"))

    def test_uses_communication_name_as_title_and_cleans_table_labels(self) -> None:
        document = Document()
        title = document.add_paragraph()
        title.alignment = 1
        title_run = title.add_run("Zadávací dokumentace pro test")
        title_run.bold = True
        title_run.font.size = 16 * 12700
        document.add_heading("2.1 Název komunikace", level=2)
        document.add_paragraph(
            "Do názvu komunikace se vyplní obchodní název."
        )
        document.add_paragraph("Testovací oznámení")
        document.add_heading("2.1.1 Email", level=3)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Tribe (kam spadá požadavek)"
        table.cell(0, 1).text = "TEST"
        table.cell(1, 0).text = "Zdrojový systém (systém, který odesílá request)"
        table.cell(1, 1).text = "TEST_APP"
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "test.docx")

        self.assertTrue(markdown.startswith("# Testovací oznámení\n"))
        self.assertNotIn("Zadávací dokumentace", markdown)
        self.assertIn("- **Tribe:** TEST", markdown)
        self.assertIn("- **Zdrojový systém:** TEST_APP", markdown)
        self.assertNotIn("kam spadá požadavek", markdown)

    def test_uses_formatted_cover_title_and_preserves_bullets(self) -> None:
        document = Document()
        title = document.add_paragraph()
        title.alignment = 1
        run = title.add_run("Název zadávací dokumentace")
        run.bold = True
        run.font.size = 16 * 12700
        document.add_heading("1 Management summary", level=1)
        document.add_heading("2 Obchodní zadání", level=1)
        document.add_heading("2.1 Přehled", level=2)
        bullet = document.add_paragraph("Český jazyk", style="List Paragraph")
        numbering = document.part.numbering_part.element
        abstract_ids = [
            int(item.get(qn("w:abstractNumId")))
            for item in numbering.xpath("./w:abstractNum")
        ]
        abstract_id = max(abstract_ids, default=-1) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        number_format = OxmlElement("w:numFmt")
        number_format.set(qn("w:val"), "bullet")
        level.append(number_format)
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), "•")
        level.append(level_text)
        abstract.append(level)
        numbering.append(abstract)
        num_id = max((item.numId for item in numbering.num_lst), default=0) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_reference = OxmlElement("w:abstractNumId")
        abstract_reference.set(qn("w:val"), str(abstract_id))
        num.append(abstract_reference)
        numbering.append(num)
        properties = bullet._p.get_or_add_pPr()
        num_properties = OxmlElement("w:numPr")
        level_reference = OxmlElement("w:ilvl")
        level_reference.set(qn("w:val"), "0")
        num_reference = OxmlElement("w:numId")
        num_reference.set(qn("w:val"), str(num_id))
        num_properties.extend([level_reference, num_reference])
        properties.append(num_properties)
        document.add_paragraph("Příklad běžného odstavce", style="List Paragraph")
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "zadani.docx")

        self.assertIn("# Název zadávací dokumentace", markdown)
        self.assertIn("## Základní informace", markdown)
        self.assertIn("## Obchodní zadání", markdown)
        self.assertIn("### Přehled", markdown)
        self.assertNotIn("Management summary", markdown)
        self.assertNotIn("## 2 Obchodní zadání", markdown)
        self.assertIn("- Český jazyk", markdown)
        self.assertIn("Příklad běžného odstavce", markdown)
        self.assertNotIn("- Příklad běžného odstavce", markdown)

    def test_formats_form_fields_and_message_text(self) -> None:
        document = Document()
        document.add_paragraph("Číslo JIRA epicu: ABC-123")
        label = document.add_paragraph()
        label.add_run("Text SMS").bold = True
        document.add_paragraph("Dobrý den, <jméno>.")
        document.add_paragraph("Vaše banka")
        document.add_paragraph("Obsah emailu:")
        document.add_paragraph("Dobrý den,")
        document.add_paragraph("Toto je text e-mailu.")
        document.add_paragraph("Preheader emailu:")
        document.add_paragraph("Krátký náhled zprávy.")
        document.add_paragraph("CS:\nTitulek zprávy\nText zprávy")
        document.add_heading("Další část", level=1)
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "komunikace.docx")

        self.assertIn("**Číslo JIRA epicu:** ABC-123", markdown)
        self.assertIn(
            "**Text SMS**\n\n```text\nDobrý den, <jméno>.\n\nVaše banka\n```",
            markdown,
        )
        self.assertIn(
            "**Tělo emailu**\n\n"
            "```text\n<div>Dobrý den,</div><div>Toto je text e-mailu.</div>\n```",
            markdown,
        )
        self.assertIn(
            "**Preheader emailu**\n\n```text\nKrátký náhled zprávy.\n```",
            markdown,
        )
        self.assertIn("```text\nCS:\nTitulek zprávy\nText zprávy\n```", markdown)

    def test_formats_email_fields_html_and_anchored_comments(self) -> None:
        document = Document()
        document.add_heading("2.1.2 Email", level=1)
        document.add_paragraph("Český jazyk", style="List Paragraph")
        document.add_paragraph("Předmět emailu: Předmět mailu")
        document.add_paragraph("Preheader emailu:")
        document.add_paragraph("Krátký náhled")
        body = document.add_paragraph()
        body.add_run("Tento text je ")
        body.add_run("tučně").bold = True
        body.add_run(" a obsahuje ")
        body.add_run("<contractNum>")
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(
            self.add_comment_to_docx(output.getvalue()),
            "email.docx",
        )

        self.assertIn("## Email\n\n### CZ varianta", markdown)
        self.assertIn(
            "**Předmět emailu**\n\n```text\nPředmět mailu\n```",
            markdown,
        )
        self.assertIn(
            "**Preheader emailu**\n\n```text\nKrátký náhled\n```",
            markdown,
        )
        self.assertIn(
            "**Tělo emailu**\n\n"
            "```text\n<div>Tento text je <b>tučně</b> a obsahuje "
            "[1]&lt;contractNum&gt;[/1]</div>\n```",
            markdown,
        )
        self.assertIn(
            "#### Komentáře k tělu e-mailu\n\n"
            "**[1] Komentář**\n\n```text\nČíslo smlouvy\n```",
            markdown,
        )

    def test_omits_email_template_guidance_but_preserves_email_content(self) -> None:
        instruction = (
            "Tato podkapitola se vyplní v případě, pokud bude v požadavku "
            "definovaná alespoň jedna E-mailová textace. U všech emailů se "
            "vkládá standardní hlavička a patička ČS."
        )
        placeholder = (
            "Zde se doplní Emailová textace, která bude v českém jazyce."
        )
        document = Document()
        document.add_heading("Email", level=1)
        document.add_paragraph(instruction)
        document.add_paragraph("Český jazyk", style="List Paragraph")
        document.add_paragraph("Předmět emailu: Oznámení o změně")
        document.add_paragraph(placeholder)
        body = document.add_paragraph()
        body.add_run("Váš účet ")
        body.add_run("<contractNum>")
        body.add_run(" byl aktualizován.")
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(
            self.add_comment_to_docx(output.getvalue()),
            "email.docx",
        )

        self.assertNotIn("Tato podkapitola se vyplní", markdown)
        self.assertNotIn(placeholder, markdown)
        self.assertIn("Oznámení o změně", markdown)
        self.assertIn("Váš účet", markdown)
        self.assertIn("Komentáře k tělu e-mailu", markdown)
        self.assertIn("Číslo smlouvy", markdown)

    def test_infers_email_body_after_subject_when_no_body_label_exists(self) -> None:
        document = Document()
        document.add_heading("Email", level=1)
        document.add_paragraph("CZ varianta", style="List Paragraph")
        document.add_paragraph("Předmět emailu: Testovací předmět")
        document.add_paragraph("První odstavec e-mailu.")
        document.add_paragraph("Druhý odstavec e-mailu.")
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "email.docx")

        self.assertIn(
            "**Tělo emailu**\n\n"
            "```text\n<div>První odstavec e-mailu.</div>"
            "<div>Druhý odstavec e-mailu.</div>\n```",
            markdown,
        )

    def add_comment_to_docx(self, content: bytes) -> bytes:
        with ZipFile(BytesIO(content)) as source:
            files = {
                item.filename: source.read(item.filename)
                for item in source.infolist()
            }

        namespace = f"{{{_WORD_NS}}}"
        root = ElementTree.fromstring(files["word/document.xml"])
        parents = {
            child: parent
            for parent in root.iter()
            for child in parent
        }
        anchor_run = next(
            run
            for run in root.iter(f"{namespace}r")
            if "".join(
                text.text or ""
                for text in run.iter(f"{namespace}t")
            ) == "<contractNum>"
        )
        paragraph = parents[anchor_run]
        index = list(paragraph).index(anchor_run)
        start = ElementTree.Element(f"{namespace}commentRangeStart")
        start.set(f"{namespace}id", "7")
        end = ElementTree.Element(f"{namespace}commentRangeEnd")
        end.set(f"{namespace}id", "7")
        paragraph.insert(index, start)
        paragraph.insert(index + 2, end)
        reference_run = ElementTree.Element(f"{namespace}r")
        reference = ElementTree.SubElement(
            reference_run,
            f"{namespace}commentReference",
        )
        reference.set(f"{namespace}id", "7")
        paragraph.insert(index + 3, reference_run)
        files["word/document.xml"] = ElementTree.tostring(
            root,
            encoding="utf-8",
            xml_declaration=True,
        )

        comments = ElementTree.Element(f"{namespace}comments")
        comment = ElementTree.SubElement(
            comments,
            f"{namespace}comment",
            {f"{namespace}id": "7", f"{namespace}author": "Tester"},
        )
        comment_paragraph = ElementTree.SubElement(comment, f"{namespace}p")
        comment_run = ElementTree.SubElement(comment_paragraph, f"{namespace}r")
        comment_text = ElementTree.SubElement(comment_run, f"{namespace}t")
        comment_text.text = "Číslo smlouvy"
        files["word/comments.xml"] = ElementTree.tostring(
            comments,
            encoding="utf-8",
            xml_declaration=True,
        )

        result = BytesIO()
        with ZipFile(result, "w", ZIP_DEFLATED) as archive:
            for filename, data in files.items():
                archive.writestr(filename, data)
        return result.getvalue()

    def test_splits_push_title_from_message_text(self) -> None:
        document = Document()
        document.add_heading("PUSH", level=1)
        document.add_paragraph("CS:\nČeský titulek\nČeský text zprávy")
        document.add_paragraph("CS:\nEnglish title\nEnglish message text")
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "push.docx")

        self.assertEqual(markdown.count("- **Segmentace:** CS"), 1)
        self.assertEqual(markdown.count("**Titulek PUSH**"), 2)
        self.assertEqual(markdown.count("**Text PUSH**"), 2)
        self.assertIn("```text\nČeský titulek\n```", markdown)
        self.assertIn("```text\nČeský text zprávy\n```", markdown)
        self.assertIn("```text\nEnglish title\n```", markdown)
        self.assertIn("```text\nEnglish message text\n```", markdown)

    def test_omits_channels_marked_not_required(self) -> None:
        cases = (
            ("Email", "Email není požadována", "after_heading"),
            ("SMS", "SMS není požadována", "before_heading"),
            ("PUSH", "PUSH není požadována", "after_body"),
        )
        for channel, status, status_position in cases:
            with self.subTest(channel=channel):
                document = Document()
                if status_position == "before_heading":
                    document.add_paragraph(status)
                document.add_heading(channel, level=1)
                document.add_paragraph(f"Obsah sekce {channel}.")
                if status_position == "after_heading":
                    document.add_paragraph(status)
                elif status_position == "after_body":
                    document.add_paragraph(status)
                output = BytesIO()
                document.save(output)

                markdown, _ = convert_docx(
                    output.getvalue(),
                    "nepozadovana-komunikace.docx",
                )

                self.assertNotIn(f"## {channel}", markdown)
                self.assertNotIn(status, markdown)
                self.assertNotIn(f"Obsah sekce {channel}.", markdown)

    def test_omits_instruction_section_until_next_peer_heading(self) -> None:
        document = Document()
        document.add_heading("Základní informace", level=1)
        document.add_paragraph("Důležitý obsah.")
        document.add_heading("3 Instrukce", level=2)
        document.add_paragraph("Instrukce k vyplnění.")
        document.add_heading("Příklad", level=3)
        document.add_paragraph("Další text instrukcí.")
        document.add_heading("Obchodní zadání", level=2)
        document.add_paragraph("Zachovaný obsah.")
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "pozadavek.docx")

        self.assertNotIn("## Instrukce", markdown)
        self.assertNotIn("Instrukce k vyplnění.", markdown)
        self.assertNotIn("Další text instrukcí.", markdown)
        self.assertIn("Základní informace", markdown)
        self.assertIn("Obchodní zadání", markdown)
        self.assertIn("Zachovaný obsah.", markdown)

    def test_omits_instruction_field_from_docx_paragraph_and_table(self) -> None:
        instruction = (
            "Cílem této kapitoly je vytvořit shrnutí obsahu požadavku na EOM. "
            "Vyplní Business garant/ Business analytik."
        )
        document = Document()
        document.add_paragraph(f"Instrukce: {instruction}")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Instrukce"
        table.rows[0].cells[1].text = instruction
        document.add_paragraph("Požadovaný obsah.")
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "pozadavek.docx")

        self.assertNotIn("Instrukce", markdown)
        self.assertNotIn("Cílem této kapitoly", markdown)
        self.assertIn("Požadovaný obsah.", markdown)

    def test_omits_optional_note_below_table(self) -> None:
        document = Document()
        document.add_table(rows=1, cols=2)
        document.add_paragraph("* nepovinné")
        document.add_paragraph("Důležitý obsah.")
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "pozadavek.docx")

        self.assertNotIn("nepovinné", markdown)
        self.assertIn("Důležitý obsah.", markdown)

    def test_formats_guarantor_details_as_template_sections(self) -> None:
        document = Document()
        table = document.add_table(rows=0, cols=2)
        rows = (
            ("Tribe", "T302_IT Corporate Daily Banking and Lending"),
            (
                "Business garant",
                "Iva Zázvorková\nemail: izazvorkova@csas.cz\n"
                "telefon: +420 704 993 082",
            ),
            (
                "Garant pro akceptaci požadavku a schválení pro nasazen "
                "na PROD prostředí",
                "Iva Zázvorková\nemail: izazvorkova@csas.cz\n"
                "telefon: +420 704 993 082",
            ),
            ("Zdrojový systém", "G4B_NOTIFICATIONS"),
        )
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "pozadavek.docx")

        self.assertIn("## Business garant", markdown)
        self.assertIn("- **Jméno:** Iva Zázvorková", markdown)
        self.assertIn("- **E-mail:** izazvorkova@csas.cz", markdown)
        self.assertIn("- **Telefon:** +420704993082", markdown)
        self.assertIn(
            "## Garant akceptace a schválení nasazení na PROD\n\n"
            "Stejný jako Business garant",
            markdown,
        )
        self.assertIn("- **Tribe:** T302_IT Corporate Daily Banking and Lending", markdown)
        self.assertIn("- **Zdrojový systém:** G4B_NOTIFICATIONS", markdown)
        self.assertNotIn("<br>email:", markdown)
        self.assertNotIn("+420 704 993 082", markdown)

    def test_formats_business_guarantor_paragraph_as_section(self) -> None:
        document = Document()
        document.add_paragraph(
            "Business garant: Iva Zázvorková<br>"
            "email: izazvorkova@csas.cz<br>"
            "telefon: +420 704 993 082"
        )
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "pozadavek.docx")

        self.assertIn(
            "## Business garant\n\n"
            "- **Jméno:** Iva Zázvorková\n"
            "- **E-mail:** izazvorkova@csas.cz\n"
            "- **Telefon:** +420704993082",
            markdown,
        )
        self.assertNotIn("<br>", markdown)

    def test_collapses_matching_business_and_acceptance_paragraphs(self) -> None:
        contact = (
            "Iva Zázvorková<br>email: izazvorkova@csas.cz<br>"
            "telefon: +420 704 993 082"
        )
        document = Document()
        document.add_paragraph(f"Business garant: {contact}")
        document.add_paragraph(
            "Garant pro akceptaci požadavku a schválení pro nasazen "
            f"na PROD prostředí: {contact}"
        )
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "pozadavek.docx")

        self.assertIn(
            "## Garant akceptace a schválení nasazení na PROD\n\n"
            "Stejný jako Business garant",
            markdown,
        )
        self.assertEqual(markdown.count("Iva Zázvorková"), 1)

    def test_simplifies_trace_storage_label(self) -> None:
        document = Document()
        row = document.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = "Ukládání komunikace do TRACE"
        row.cells[1].text = "NE"
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "pozadavek.docx")

        self.assertIn("- **Ukládání do TRACE:** NE", markdown)
        self.assertNotIn("Ukládání komunikace do TRACE", markdown)

    def test_trace_storage_keeps_only_standalone_yes_or_no(self) -> None:
        document = Document()
        table = document.add_table(rows=0, cols=2)
        rows = (
            (
                "Ukládání komunikace do TRACE",
                "NE\nTrace_Type: typ\nTrace_Sub_Type: podtyp\nState: aktivní",
            ),
            (
                "Ukládání do TRACE",
                "Trace_Type: typ\nTrace_Sub_Type: podtyp\nState: NE",
            ),
        )
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "pozadavek.docx")

        self.assertEqual(markdown.count("- **Ukládání do TRACE:** NE"), 1)
        self.assertNotIn("Trace_Type", markdown)
        self.assertNotIn("Trace_Sub_Type", markdown)
        self.assertNotIn("State:", markdown)

    def test_normalizes_kcp_consultation_to_yes_no_and_defaults_blank(self) -> None:
        document = Document()
        table = document.add_table(rows=0, cols=2)
        rows = (
            (
                "Požadavek musí být konzultován s KCP. E-mailová adresa, "
                "na kterou se obraťte: pripominkyWFM@csas.cz",
                "ANO",
            ),
            ("Konzultace s KCP", ""),
        )
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value
        output = BytesIO()
        document.save(output)

        markdown, _ = convert_docx(output.getvalue(), "pozadavek.docx")

        self.assertIn("- **Konzultace s KCP:** ANO", markdown)
        self.assertIn("- **Konzultace s KCP:** NE", markdown)
        self.assertNotIn("Požadavek musí být konzultován", markdown)
        self.assertNotIn("pripominkyWFM@csas.cz", markdown)

    def test_extracts_word_comment_and_its_anchor_text(self) -> None:
        document_xml = f"""\
<w:document xmlns:w="{_WORD_NS}">
  <w:body><w:p>
    <w:commentRangeStart w:id="7"/>
    <w:r><w:t>licenceName</w:t></w:r>
    <w:commentRangeEnd w:id="7"/>
  </w:p></w:body>
</w:document>""".encode()
        comments_xml = f"""\
<w:comments xmlns:w="{_WORD_NS}">
  <w:comment w:id="7" w:author="Tester"><w:p><w:r><w:t>Proměnná</w:t></w:r></w:p></w:comment>
</w:comments>""".encode()

        comments = _extract_comments(document_xml, comments_xml)

        self.assertEqual(len(comments), 1)
        self.assertEqual(comments[0].text, "Proměnná")
        self.assertEqual(comments[0].quoted_text, "licenceName")


class ConvertOtherFormatsTests(unittest.TestCase):
    def test_converts_utf8_text_and_adds_fallback_title(self) -> None:
        markdown, warnings = convert_txt(
            b"\xef\xbb\xbfPrvn\xc3\xad \xc5\x99\xc3\xa1dek.\nDruh\xc3\xbd \xc5\x99\xc3\xa1dek.",
            "poznamky.txt",
        )

        self.assertEqual(markdown, "# poznamky\n\nPrvní řádek.\nDruhý řádek.\n")
        self.assertEqual(warnings, [])

    def test_omits_template_guidance_from_text_but_keeps_content(self) -> None:
        content = (
            "Nadpis\n\n"
            "Tato podkapitola se vyplní v případě definované textace.\n\n"
            "Zde se doplní Emailová textace, která bude v českém jazyce.\n\n"
            "Váš účet je aktivní."
        )

        markdown, _ = convert_txt(content.encode("utf-8"), "poznamky.txt")

        self.assertNotIn("Tato podkapitola se vyplní", markdown)
        self.assertNotIn("Zde se doplní Emailová textace", markdown)
        self.assertIn("Váš účet je aktivní.", markdown)

    def test_omits_markdown_instruction_field_from_text(self) -> None:
        content = (
            "## Základní informace\n\n"
            "- **Instrukce:** Cílem této kapitoly je vytvořit shrnutí obsahu "
            "požadavku na EOM. Vyplní Business garant/ Business analytik.\n\n"
            "- **Tribe:** DAF"
        )

        markdown, _ = convert_txt(content.encode("utf-8"), "pozadavek.txt")

        self.assertNotIn("Instrukce", markdown)
        self.assertNotIn("Cílem této kapitoly", markdown)
        self.assertIn("- **Tribe:** DAF", markdown)

    def test_omits_standalone_optional_marker_from_text(self) -> None:
        content = "Hodnota\n\n*nepovinné\n\nDalší obsah."

        markdown, _ = convert_txt(content.encode("utf-8"), "pozadavek.txt")

        self.assertNotIn("*nepovinné", markdown)
        self.assertIn("Hodnota", markdown)
        self.assertIn("Další obsah.", markdown)

    def test_rejects_text_that_is_not_utf8(self) -> None:
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            convert_txt(b"\xff", "poznamky.txt")

    def test_converts_pdf_without_extractable_text_with_warning(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        output = BytesIO()
        writer.write(output)

        markdown, warnings = convert_pdf(output.getvalue(), "sken.pdf")

        self.assertEqual(markdown, "# sken\n")
        self.assertTrue(any("OCR" in warning for warning in warnings))
        self.assertTrue(any("nebyl nalezen" in warning for warning in warnings))

    def test_rejects_invalid_pdf(self) -> None:
        with self.assertRaisesRegex(ValueError, "platný nebo čitelný"):
            convert_pdf(b"not a PDF", "vadny.pdf")


_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


if __name__ == "__main__":
    unittest.main()
